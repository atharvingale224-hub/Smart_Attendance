import math
import time
import secrets
import io
import csv
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from models import db, College, GeofenceConfig, AcademicClass, StudentProfile, LectureSession, AttendanceRecord, ProxyRiskLog

class GeoService:
    @staticmethod
    def calculate_distance(lat1, lon1, lat2, lon2):
        """Haversine distance formula between two GPS points in meters"""
        R = 6371000.0  # Earth radius in meters
        phi1 = math.radians(lat1)
        phi2 = math.radians(lat2)
        delta_phi = math.radians(lat2 - lat1)
        delta_lambda = math.radians(lon2 - lon1)

        a = math.sin(delta_phi / 2.0)**2 + \
            math.cos(phi1) * math.cos(phi2) * \
            math.sin(delta_lambda / 2.0)**2

        c = 2.0 * math.atan2(math.sqrt(a), math.sqrt(1.0 - a))
        return R * c

    @classmethod
    def validate_location(cls, geofence, student_lat, student_lng, accuracy):
        """
        Server-side GPS verification:
        1. Validate coordinates range
        2. Validate reported GPS accuracy (reject poor/mocked accuracy)
        3. Calculate distance from campus center
        4. Validate against configured geofence radius
        """
        if not (-90.0 <= student_lat <= 90.0) or not (-180.0 <= student_lng <= 180.0):
            return False, 0.0, "Invalid GPS latitude/longitude values received."
            
        if accuracy is not None and accuracy > geofence.max_accuracy_meters:
            return False, 0.0, f"GPS accuracy too poor ({round(accuracy, 1)}m). Must be within {geofence.max_accuracy_meters}m. Please move near a window or check location settings."

        distance = cls.calculate_distance(
            geofence.latitude, 
            geofence.longitude, 
            student_lat, 
            student_lng
        )

        if distance > geofence.radius_meters:
            return False, distance, f"Outside campus geofence boundary. Distance: {round(distance, 1)}m (Allowed: {geofence.radius_meters}m)."

        return True, distance, "Location verified within campus geofence."

class QRTokenService:
    @staticmethod
    def generate_token():
        """Generates a cryptographically secure random token (256-bit URL-safe)"""
        return secrets.token_urlsafe(32)

    @classmethod
    def get_or_rotate_token(cls, session, max_age_seconds=80):
        """
        If current token is older than max_age_seconds, generates a new rotating token.
        Returns (current_token, time_remaining_seconds)
        """
        now = time.time()
        elapsed = now - session.token_generated_at
        
        if elapsed >= max_age_seconds:
            session.current_token = cls.generate_token()
            session.token_generated_at = now
            session.token_sequence = (session.token_sequence or 0) + 1
            db.session.commit()
            elapsed = 0.0

        remaining = max(0.0, max_age_seconds - elapsed)
        return session.current_token, int(remaining)

    @classmethod
    def verify_token(cls, session, submitted_token, max_age_seconds=80):
        """
        Strict server-side token validation:
        1. Session must be ACTIVE
        2. Token string must match current token
        3. Server elapsed time must be <= max_age_seconds
        """
        if session.status != 'ACTIVE':
            return False, "Attendance session has already been closed."

        if session.current_token != submitted_token:
            return False, "Invalid or outdated QR token. Please scan the current rotating QR code on the projector."

        elapsed = time.time() - session.token_generated_at
        if elapsed > max_age_seconds:
            return False, f"QR code token expired ({round(elapsed, 1)}s elapsed, maximum {max_age_seconds}s). Please scan the freshly updated code."

        return True, "Token valid."

class AttendanceService:
    @staticmethod
    def mark_attendance(
        session_id, 
        prn, 
        latitude, 
        longitude, 
        accuracy=10.0, 
        device_uuid=None, 
        ip_address=None, 
        user_agent=None,
        submitted_token=None
    ):
        session = db.session.get(LectureSession, session_id)
        if not session:
            return {'status': 'failed', 'message': 'Attendance session not found.'}, 404

        if session.status != 'ACTIVE':
            ProxyRiskLog_entry = ProxyRiskLog(
                college_id=session.college_id,
                session_id=session.id,
                prn=prn,
                risk_type='UNAUTHORIZED_ACCESS',
                details=f"Attempted to mark attendance on closed session {session.id}",
                ip_address=ip_address,
                user_agent=user_agent
            )
            db.session.add(ProxyRiskLog_entry)
            db.session.commit()
            return {'status': 'failed', 'message': 'This attendance session has already been closed.'}, 400

        # Verify QR Token
        if submitted_token:
            token_valid, token_msg = QRTokenService.verify_token(session, submitted_token)
            if not token_valid:
                ProxyRiskLog_entry = ProxyRiskLog(
                    college_id=session.college_id,
                    session_id=session.id,
                    prn=prn,
                    risk_type='EXPIRED_TOKEN',
                    details=f"Token verification failed: {token_msg}",
                    ip_address=ip_address,
                    user_agent=user_agent
                )
                db.session.add(ProxyRiskLog_entry)
                db.session.commit()
                return {'status': 'failed', 'message': token_msg}, 400

        # Look up Student
        student = StudentProfile.query.filter_by(
            college_id=session.college_id, 
            prn=prn
        ).first()

        if not student:
            return {'status': 'failed', 'message': f'Student with PRN {prn} is not registered in this college.'}, 404

        # Validate Student belongs to the lecture's class
        if student.class_id != session.class_id:
            return {'status': 'failed', 'message': f'Student PRN {prn} does not belong to this class ({session.academic_class.name} Div {session.academic_class.division}).'}, 403

        # Device Binding Validation
        if device_uuid:
            if not student.registered_device_uuid:
                # First-time registration of student's device
                student.registered_device_uuid = device_uuid
                student.device_registered_at = datetime.utcnow()
                db.session.commit()
            elif student.registered_device_uuid != device_uuid:
                risk_entry = ProxyRiskLog(
                    college_id=session.college_id,
                    session_id=session.id,
                    student_id=student.id,
                    prn=prn,
                    risk_type='DEVICE_MISMATCH',
                    details=f"Device mismatch! Registered: {student.registered_device_uuid[:8]}..., Submitted: {device_uuid[:8]}...",
                    ip_address=ip_address,
                    user_agent=user_agent
                )
                db.session.add(risk_entry)
                db.session.commit()
                return {
                    'status': 'failed', 
                    'message': 'Account is bound to a different mobile device. Please contact your college administrator to request a device reset.'
                }, 403

        # Geofencing Validation
        college = db.session.get(College, session.college_id)
        geofence = college.geofence if college else None
        
        if not geofence:
            # Fallback default geofence if not configured yet
            geofence = GeofenceConfig(
                college_id=session.college_id,
                latitude=16.6800,
                longitude=74.2300,
                radius_meters=150.0,
                max_accuracy_meters=50.0
            )

        geo_valid, distance, geo_msg = GeoService.validate_location(
            geofence, 
            latitude, 
            longitude, 
            accuracy
        )

        if not geo_valid:
            risk_entry = ProxyRiskLog(
                college_id=session.college_id,
                session_id=session.id,
                student_id=student.id,
                prn=prn,
                risk_type='OUT_OF_BOUNDS',
                details=f"Geofence violation: {geo_msg} (Lat: {latitude}, Lng: {longitude}, Acc: {accuracy}m)",
                ip_address=ip_address,
                user_agent=user_agent
            )
            db.session.add(risk_entry)
            db.session.commit()
            return {'status': 'failed', 'message': geo_msg}, 400

        # Duplicate Attendance Check
        existing_record = AttendanceRecord.query.filter_by(
            session_id=session.id, 
            student_id=student.id
        ).first()

        if existing_record:
            risk_entry = ProxyRiskLog(
                college_id=session.college_id,
                session_id=session.id,
                student_id=student.id,
                prn=prn,
                risk_type='DUPLICATE_ATTEMPT',
                details=f"Duplicate attendance submission attempt by student PRN {prn}",
                ip_address=ip_address,
                user_agent=user_agent
            )
            db.session.add(risk_entry)
            db.session.commit()
            return {'status': 'failed', 'message': f'Attendance has already been marked for PRN {prn} in this lecture!'}, 400

        # Atomic Transaction Attendance Creation
        try:
            record = AttendanceRecord(
                college_id=session.college_id,
                session_id=session.id,
                student_id=student.id,
                prn=student.prn,
                distance_meters=distance,
                accuracy_meters=accuracy or 10.0,
                status='PRESENT',
                timestamp=datetime.utcnow()
            )
            db.session.add(record)
            db.session.commit()
            
            return {
                'status': 'success',
                'message': f'Attendance successfully marked for {student.user.full_name} (PRN: {prn})!',
                'record': record.to_dict()
            }, 200
        except Exception as e:
            db.session.rollback()
            return {'status': 'failed', 'message': f'Failed to record attendance: {str(e)}'}, 500

class ReportService:
    @staticmethod
    def get_session_attendance(session_id):
        session = db.session.get(LectureSession, session_id)
        if not session:
            return None

        # All enrolled students in class
        enrolled_students = StudentProfile.query.filter_by(class_id=session.class_id).all()
        attendance_records = AttendanceRecord.query.filter_by(session_id=session_id).all()
        
        present_student_ids = {r.student_id for r in attendance_records}
        
        present_list = [r.to_dict() for r in attendance_records]
        absent_list = [
            s.to_dict() for s in enrolled_students 
            if s.id not in present_student_ids
        ]
        
        total_enrolled = len(enrolled_students)
        total_present = len(present_list)
        percentage = round((total_present / total_enrolled * 100.0), 2) if total_enrolled > 0 else 0.0

        return {
            'session': session.to_dict(),
            'total_enrolled': total_enrolled,
            'total_present': total_present,
            'total_absent': len(absent_list),
            'attendance_percentage': percentage,
            'present_students': present_list,
            'absent_students': absent_list
        }

    @staticmethod
    def get_defaulter_list(class_id, threshold_percentage=75.0):
        students = StudentProfile.query.filter_by(class_id=class_id).all()
        # All conducted sessions for this class
        conducted_sessions = LectureSession.query.filter_by(class_id=class_id, status='CLOSED').all()
        total_conducted = len(conducted_sessions)
        
        if total_conducted == 0:
            return {
                'total_conducted_lectures': 0,
                'threshold_percentage': threshold_percentage,
                'defaulters': []
            }

        defaulters = []
        for student in students:
            present_count = AttendanceRecord.query.filter(
                AttendanceRecord.student_id == student.id,
                AttendanceRecord.session_id.in_([s.id for s in conducted_sessions])
            ).count()

            percentage = round((present_count / total_conducted * 100.0), 2)
            if percentage < threshold_percentage:
                defaulters.append({
                    'student_id': student.id,
                    'prn': student.prn,
                    'roll_no': student.roll_no,
                    'name': student.user.full_name if student.user else None,
                    'email': student.user.email if student.user else None,
                    'present_lectures': present_count,
                    'conducted_lectures': total_conducted,
                    'attendance_percentage': percentage
                })

        return {
            'total_conducted_lectures': total_conducted,
            'threshold_percentage': threshold_percentage,
            'defaulters': defaulters
        }

    @staticmethod
    def export_session_csv(session_id):
        data = ReportService.get_session_attendance(session_id)
        if not data:
            return None
            
        output = io.StringIO()
        writer = csv.writer(output)
        
        writer.writerow(["ATTENDIO ATTENDANCE REPORT"])
        writer.writerow(["Session ID", data['session']['id']])
        writer.writerow(["Subject", f"{data['session']['subject_name']} ({data['session']['subject_code']})"])
        writer.writerow(["Class", f"{data['session']['class_name']} Div {data['session']['division']}"])
        writer.writerow(["Teacher", data['session']['teacher_name']])
        writer.writerow(["Started At", data['session']['started_at']])
        writer.writerow(["Total Students", data['total_enrolled']])
        writer.writerow(["Present Count", data['total_present']])
        writer.writerow(["Percentage", f"{data['attendance_percentage']}%"])
        writer.writerow([])
        writer.writerow(["Status", "PRN", "Roll No", "Student Name", "Time", "Distance (m)", "Accuracy (m)"])
        
        for p in data['present_students']:
            writer.writerow(["PRESENT", p['prn'], p.get('roll_no', ''), p.get('student_name', ''), p['timestamp'], p['distance_meters'], p['accuracy_meters']])
            
        for a in data['absent_students']:
            writer.writerow(["ABSENT", a['prn'], a.get('roll_no', ''), a.get('full_name', ''), "-", "-", "-"])
            
        return output.getvalue()

    @staticmethod
    def export_session_excel(session_id):
        data = ReportService.get_session_attendance(session_id)
        if not data:
            return None
            
        wb = Workbook()
        ws = wb.active
        ws.title = "Attendance Summary"
        
        header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        
        ws.append(["ATTENDIO ATTENDANCE REPORT"])
        ws.append(["Subject", f"{data['session']['subject_name']} ({data['session']['subject_code']})"])
        ws.append(["Class", f"{data['session']['class_name']} Div {data['session']['division']}"])
        ws.append(["Teacher", data['session']['teacher_name']])
        ws.append(["Attendance %", f"{data['attendance_percentage']}%"])
        ws.append([])
        
        table_headers = ["Status", "PRN", "Roll No", "Student Name", "Time", "Distance (m)", "Accuracy (m)"]
        ws.append(table_headers)
        
        header_row = ws.max_row
        for col_num in range(1, len(table_headers) + 1):
            cell = ws.cell(row=header_row, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
            
        for p in data['present_students']:
            ws.append(["PRESENT", p['prn'], p.get('roll_no', ''), p.get('student_name', ''), p['timestamp'], p['distance_meters'], p['accuracy_meters']])
            
        for a in data['absent_students']:
            ws.append(["ABSENT", a['prn'], a.get('roll_no', ''), a.get('full_name', ''), "-", "-", "-"])
            
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()
